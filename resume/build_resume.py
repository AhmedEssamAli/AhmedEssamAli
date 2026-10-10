#!/usr/bin/env python3
"""Build the ATS resume (DOCX + PDF) from profile/career.yaml."""
import re
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "profile" / "career.yaml"
OUT = HERE / "output"
BASENAME = "Ahmed_Essam_Ali_Resume"

FONT = "Calibri"
MARGIN = 0.65
PAGE_W = 8.5
TAB_POS = Inches(PAGE_W - 2 * MARGIN)

# Spacing knobs (points), tuned for exactly 2 pages.
BULLET_AFTER = 2
ENTRY_BEFORE = 6
HEAD_BEFORE = 10
LINE_SPACING = 1.0


def set_font(style_or_run, size=None, bold=None):
    f = style_or_run.font
    f.name = FONT
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") and hasattr(style_or_run.element, "get_or_add_rPr") else None
    if rpr is None:
        rpr = style_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)


def fmt(p, before=0, after=0, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = LINE_SPACING
    if align is not None:
        p.alignment = align


def add_run(p, text, size=10.5, bold=False):
    r = p.add_run(text)
    set_font(r, size, bold)
    return r


def bottom_border(p):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "000000")
    bdr.append(b)
    ppr.append(bdr)


def heading(doc, text):
    p = doc.add_paragraph()
    fmt(p, before=HEAD_BEFORE, after=3)
    p.paragraph_format.keep_with_next = True
    add_run(p, text.upper(), 11.5, True)
    bottom_border(p)


def line_with_date(doc, left, date, bold_left=False, before=0, after=0, keep=False):
    p = doc.add_paragraph()
    fmt(p, before=before, after=after)
    p.paragraph_format.tab_stops.add_tab_stop(TAB_POS, WD_TAB_ALIGNMENT.RIGHT)
    if keep:
        p.paragraph_format.keep_with_next = True
    add_run(p, left, 10.5, bold_left)
    if date:
        add_run(p, "\t" + date, 10.5, False)
    return p


def build_docx(d, path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(MARGIN)
    sec.top_margin = sec.bottom_margin = Inches(0.6)

    set_font(doc.styles["Normal"], 10.5)

    cp = doc.core_properties
    cp.title = "Ahmed Essam Ali Resume"
    cp.author = "Ahmed Essam Ali"

    c = d["contact"]
    p = doc.add_paragraph(); fmt(p, 0, 2, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, d["name"], 18, True)
    p = doc.add_paragraph(); fmt(p, 0, 2, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, d["headline"])
    p = doc.add_paragraph(); fmt(p, 0, 2, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, " | ".join([c["location"], c["phone"], c["email"], c["linkedin"]]))
    p = doc.add_paragraph(); fmt(p, 0, 0, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, c["relocation"])

    heading(doc, "Summary")
    p = doc.add_paragraph(); fmt(p)
    add_run(p, " ".join(d["summary"].split()))

    heading(doc, "Experience")
    for i, e in enumerate(d["experience"]):
        left = f'{e["title"]} | {e["company"]} | {e["location"]}'
        dates = f'{e["start"]} – {e["end"]}'
        line_with_date(doc, left, dates, True, before=0 if i == 0 else ENTRY_BEFORE,
                       after=1.5, keep=True)
        for b in e["bullets"]:
            # Plain "•" text, not the List Bullet style: its Symbol-font glyph
            # extracts from the PDF as private-use U+F0B7.
            bp = doc.add_paragraph()
            fmt(bp, 0, BULLET_AFTER)
            bp.paragraph_format.left_indent = Inches(0.25)
            bp.paragraph_format.first_line_indent = Inches(-0.18)
            bp.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))
            add_run(bp, "•\t" + b)

    heading(doc, "Education")
    for e in d["education"]:
        school = re.sub(r"\s*\(.*?\)", "", e["school"])  # keep the line on one row
        line_with_date(doc, f'{e["degree"]}, {school}', str(e.get("date") or "").strip(), after=1.5)

    heading(doc, "Certifications")
    for e in d["certifications"]:
        line_with_date(doc, e["name"], str(e.get("date") or "").strip(), after=1.5)

    heading(doc, "Skills")
    for k, v in d["skills"].items():
        p = doc.add_paragraph(); fmt(p, 0, 1.5)
        add_run(p, f"{k.capitalize()}: ", 10.5, True)
        add_run(p, ", ".join(v))

    heading(doc, "Languages")
    p = doc.add_paragraph(); fmt(p)
    add_run(p, ", ".join(d["languages"]))

    doc.save(path)


def main():
    # Optional args: a tailored copy of career.yaml and an output dir (used per job).
    data = Path(sys.argv[1]) if len(sys.argv) > 1 else DATA
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else OUT
    d = yaml.safe_load(data.read_text(encoding="utf-8"))
    d.pop("needs_confirmation", None)
    out.mkdir(parents=True, exist_ok=True)
    docx = out / f"{BASENAME}.docx"
    pdf = out / f"{BASENAME}.pdf"
    build_docx(d, docx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pages = len(PdfReader(str(pdf)).pages)
    print(docx)
    print(pdf)
    print(f"Pages: {pages}")


if __name__ == "__main__":
    main()
